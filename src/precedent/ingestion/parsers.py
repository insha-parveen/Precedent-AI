"""Text extraction from contract files (PDF, DOCX, TXT)."""
from __future__ import annotations

import fitz  # pymupdf
from docx import Document as DocxDocument
from pathlib import Path
from typing import Optional


def extract_text_from_pdf(filepath: Path) -> str:
    """Extract text from a PDF file using PyMuPDF (fitz).

    Args:
        filepath: Path to the PDF file

    Returns:
        Extracted text content
    """
    text_parts = []
    with fitz.open(filepath) as doc:
        for page in doc:
            text = page.get_text("text")
            if text.strip():
                text_parts.append(text)
    return "\n\n".join(text_parts)


def extract_text_from_docx(filepath: Path) -> str:
    """Extract text from a DOCX file using python-docx.

    Args:
        filepath: Path to the DOCX file

    Returns:
        Extracted text content
    """
    doc = DocxDocument(filepath)
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_parts.append(" | ".join(row_text))
    return "\n\n".join(text_parts)


def extract_text_from_txt(filepath: Path) -> str:
    """Extract text from a plain text file.

    Args:
        filepath: Path to the TXT file

    Returns:
        File content as string
    """
    return filepath.read_text(encoding="utf-8", errors="replace")


def extract_text(filepath: Path) -> str:
    """Extract text from a file based on its extension.

    Args:
        filepath: Path to the file

    Returns:
        Extracted text content

    Raises:
        ValueError: If file extension is not supported
    """
    suffix = filepath.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(filepath)
    elif suffix == ".docx":
        return extract_text_from_docx(filepath)
    elif suffix == ".txt":
        return extract_text_from_txt(filepath)
    else:
        raise ValueError(f"Unsupported file type: {suffix}. Supported: .pdf, .docx, .txt")


def discover_contract_files(input_dir: Path, recursive: bool = True) -> list[Path]:
    """Discover all supported contract files in a directory.

    Args:
        input_dir: Directory to search
        recursive: Whether to search recursively

    Returns:
        List of file paths sorted by name
    """
    patterns = ["*.pdf", "*.docx", "*.txt"]
    files = []
    for pattern in patterns:
        if recursive:
            files.extend(input_dir.rglob(pattern))
        else:
            files.extend(input_dir.glob(pattern))
    return sorted(files, key=lambda p: str(p).lower())